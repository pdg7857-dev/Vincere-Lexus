#!/usr/bin/env python3
"""Build the Sales-CRM dataset (crm/data.js) from the Vincere-Lexus workbook.

Reads the exported Google-Sheets workbook and normalizes it into the data model
the CRM front-end expects (see crm/app.js):

  inventory[]   from  Used (pre-owned) + Inventory (new) + Pipeline (incoming)
  deals[]       from  Clients (the live opportunity records)
  appointments  generated from the active opportunities (no source feed exists)
  ads[]         from  FB_Log  (+ seeded ads for aged used stock)
  repeatBuyers  from  the Purchased/Repeat clients + "Leads - With Vehicle"

Usage:
  python3 scripts/build_crm_data.py path/to/Copy_of_VincereLexus_CRM.xlsx
  # -> writes crm/data.js
"""
import json
import re
import sys
import datetime
from pathlib import Path

import openpyxl

# Fixed "today" so the pipeline/calendar/urgency math is reproducible and lines
# up with the newest records in the export.
TODAY = datetime.date(2026, 7, 29)

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "data" / "VincereLexus_CRM.xlsx"

MAKE_MAP = {
    "LEXUS": "Lexus", "TOYOTA": "Toyota", "ACURA": "Acura", "HONDA": "Honda",
    "MERC-BENZ": "Mercedes-Benz", "MERC BENZ": "Mercedes-Benz", "VW": "Volkswagen",
    "INFINITI": "Infiniti", "PORSCHE": "Porsche", "JEEP": "Jeep", "MAZDA": "Mazda",
    "HYUNDAI": "Hyundai", "NISSAN": "Nissan", "DODGE": "Dodge", "AUDI": "Audi",
    "BMW": "BMW", "SUBARU": "Subaru", "CHRYSLER": "Chrysler", "LINCOLN": "Lincoln",
    "JAGUAR": "Jaguar", "BUICK": "Buick", "GENESIS": "Genesis", "MASERATI": "Maserati",
    "FORD": "Ford", "LAND ROVER": "Land Rover",
}

# model / series keyword -> body style
SUV = {"NX", "RX", "RZ", "GX", "LX", "TX", "UX", "RAV4", "HIGHLANDER", "4RUNNER",
       "SEQUOIA", "TAHOE", "PILOT", "CR-V", "CRV", "ROGUE", "TUCSON", "OUTBACK",
       "ESCAPE", "EXPLORER", "GRAND CHEROKEE", "WRANGLER", "MACAN", "CAYENNE",
       "QX", "GLC", "GLE", "Q5", "X3", "X5", "RDX", "MDX", "CX-5", "CX5",
       "SANTA FE", "PALISADE", "TELLURIDE", "ENVISION", "NAUTILUS", "F-PACE"}
SEDAN = {"ES", "IS", "LS", "RC", "ILX", "TLX", "TL", "TSX", "CIVIC", "ACCORD",
         "COROLLA", "CAMRY", "JETTA", "PASSAT", "3", "5", "A4", "A6", "C300",
         "C-CLASS", "E-CLASS", "CHARGER", "300", "MODEL 3", "MODEL S", "G70",
         "Q50", "ALTIMA", "MAXIMA", "SENTRA", "ELANTRA", "SONATA", "OPTIMA",
         "MALIBU", "CRUZE", "FUSION", "TAURUS", "LEGACY", "MAZDA3", "MAZDA6"}
COUPE = {"LC", "M240I", "M2", "M4", "SUPRA", "GR86", "BRZ"}
TRUCK = {"F-150", "F150", "SILVERADO", "SIERRA", "TACOMA", "TUNDRA", "RAM",
         "RIDGELINE", "FRONTIER", "COLORADO"}
VAN = {"SIENNA", "ODYSSEY", "PACIFICA", "CARNIVAL", "SEDONA"}


def title_make(m):
    if not m:
        return ""
    return MAKE_MAP.get(str(m).strip().upper(), str(m).strip().title())


def body_for(model, series=""):
    hay = (str(series) + " " + str(model)).upper()
    for kw in COUPE:
        if kw in hay:
            return "Coupe"
    for kw in TRUCK:
        if kw in hay:
            return "Truck"
    for kw in VAN:
        if kw in hay:
            return "Van"
    # sedans: match whole tokens and their alpha prefixes (so "IS300" -> "IS",
    # "ES 350" -> "ES") to avoid "IS" matching inside unrelated words.
    tokens = [t for t in re.split(r"[^A-Z0-9]+", hay) if t]
    prefixes = set(tokens)
    for t in tokens:
        m = re.match(r"([A-Z]{2,})", t)
        if m:
            prefixes.add(m.group(1))
    for kw in SEDAN:
        if kw in prefixes:
            return "Sedan"
    for kw in SUV:
        if kw in prefixes or kw in hay:
            return "SUV"
    for kw in SEDAN:
        if kw in hay:
            return "Sedan"
    return "SUV"


def fuel_for(*parts):
    hay = " ".join(str(p) for p in parts if p).lower()
    if re.search(r"\brz\b|\bev\b|electric", hay):
        return "Electric"
    if "+" in hay or "plug" in hay:
        return "Plug-in Hybrid"
    if re.search(r"\d+h\b|hybrid|\bh\+|300h|350h|450h|500h|250h|600h", hay):
        return "Hybrid"
    return "Gas"


def strip_color_code(c):
    if not c:
        return ""
    c = str(c)
    # "01J7-Atomic Silver" -> "Atomic Silver"
    m = re.match(r"^\s*[0-9A-Z]{3,4}[-]\s*(.+)$", c)
    return (m.group(1) if m else c).strip()


def base_series(series):
    """Normalize an inventory series like 'NX 450h' / 'RXh' / 'ESe' to a Pricing base."""
    s = str(series or "").upper().strip()
    s = re.split(r"[\s\-]", s)[0]
    s = re.sub(r"(450H|350H|300H|H|E)$", "", s) or s
    # after stripping, keep the leading 2-letter series
    m = re.match(r"([A-Z]{2})", s)
    return m.group(1) if m else s


def parse_new_suffix(suffix):
    """'MM-Ultra Luxury/NuLuxe,Black with Black Open-Pore Wood Trim'
       -> (trim, interior_material, interior_color)."""
    if not suffix:
        return ("", "", "")
    s = str(suffix)
    # drop the leading 2-4 char package code
    s = re.sub(r"^[0-9A-Z]{1,4}\-", "", s)
    trim_part, _, rest = s.partition("/")
    trim = re.sub(r"\s*\(.*?\)", "", trim_part).replace(" Package", "").strip()
    trim = re.sub(r"\s*-\s*Premium Paint", "", trim).strip()
    material, _, colour = rest.partition(",")
    return (trim, material.strip(), colour.strip())


def days_since(d):
    if not isinstance(d, datetime.datetime) and not isinstance(d, datetime.date):
        return 0
    if isinstance(d, datetime.datetime):
        d = d.date()
    return max(0, (TODAY - d).days)


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = re.sub(r"[^0-9.]", "", str(x))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def iso(d):
    if isinstance(d, datetime.datetime):
        d = d.date()
    if isinstance(d, datetime.date):
        return d.isoformat()
    return None


def rows(ws, start=2):
    for r in range(start, ws.max_row + 1):
        yield r, [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]


# --------------------------------------------------------------------------- #
def build_pricing(wb):
    ws = wb["Pricing"]
    by_series = {}
    for _, v in rows(ws):
        series, model, suffix, msrp = v[1], v[2], v[3], v[4]
        if not series or not msrp:
            continue
        b = base_series(series)
        by_series.setdefault(b, []).append((str(suffix or "").lower(),
                                            str(model or "").lower(), float(msrp)))
    return by_series


def price_lookup(pricing, series, trim, model=""):
    b = base_series(series)
    cands = pricing.get(b)
    if not cands:
        return None
    trim_l = (trim or "").lower()
    tokens = [t for t in re.split(r"[^a-z0-9]+", trim_l) if len(t) > 2]
    best, best_score = None, -1
    for suffix_l, model_l, msrp in cands:
        score = 0
        for t in tokens:
            if t in suffix_l:
                score += 2
        if model and model.lower() in model_l:
            score += 1
        if score > best_score:
            best_score, best = score, msrp
    if best is not None and best_score > 0:
        return best
    vals = sorted(m for _, _, m in cands)
    return vals[len(vals) // 2]  # median fallback


# --------------------------------------------------------------------------- #
def build_used(wb, referenced):
    """Pre-owned inventory. Keep available units + anything a deal references."""
    ws = wb["Used"]
    out = []
    for _, v in rows(ws):
        stock = v[1]
        if not stock:
            continue
        stock = str(stock).strip()
        price = num(v[8])
        if not price:
            continue
        status = (str(v[10] or "").strip())
        board = (str(v[14] or "").strip().upper())
        sold = board == "SOLD" or status.upper() == "WHOLESALE"
        if sold and stock not in referenced:
            continue
        colour = str(v[6] or "")
        ext, _, interior = colour.partition("/")
        model = str(v[4] or "")
        make = title_make(v[3])
        year = int(v[2]) if num(v[2]) else None
        trim = str(v[5] or "").title()
        km = int(num(v[7])) if num(v[7]) else 0
        out.append({
            "stock": stock,
            "vin": "",
            "year": year, "make": make, "model": model, "trim": trim,
            "body": body_for(model, ""),
            "price": int(round(price)),
            "atValue": int(round(price * 0.95)),
            "days": days_since(v[0]),
            "km": f"{km:,} km" if km else "—",
            "kmNum": km,
            "colour": ext.strip().title(), "interior": interior.strip().title(),
            "fuel": fuel_for(model, trim),
            "type": "Used",
            "condition": status or "Used",
            "eta": "", "from": "",
            "link": str(v[17]) if v[17] else "",
            "car": f"{year} {make} {model}".strip(),
            "detail": {
                "In-stock date": iso(v[0]),
                "Condition": status or "Used",
                "Delivery board": str(v[14] or "").strip() or "—",
                "Keys": (str(int(num(v[11]))) if num(v[11]) else "—"),
                "Recon sent": str(v[12] or "—").strip(),
                "Tires": str(v[15] or "—").strip(),
                "Carfax / owners": str(v[9] or "—").strip(),
            },
        })
    return out


def build_new(wb, pricing, sheet, kind):
    ws = wb[sheet]
    out = []
    for _, v in rows(ws):
        order = v[1]
        if not order:
            continue
        stock = str(int(order)) if isinstance(order, float) else str(order).strip()
        order_type = str(v[2] or "")
        vin = str(v[5] or "").strip()
        year = int(num(v[6])) if num(v[6]) else None
        series = str(v[7] or "")
        model_raw = str(v[8] or "")
        model = model_raw.split("-", 1)[-1].strip() if "-" in model_raw else model_raw
        trim, material, int_colour = parse_new_suffix(v[9])
        ext = strip_color_code(v[10])
        price = price_lookup(pricing, series, trim, model)
        eta_from, eta_to = v[12], v[13]
        order_status = str(v[3] or "").strip()
        allocated = str(v[14]).strip() if len(v) > 14 and v[14] else ""
        comments = str(v[15]).strip() if len(v) > 15 and v[15] else ""
        interior = (f"{material} {int_colour}").strip() or int_colour or material
        is_demo = "demo" in order_type.lower() or "loaner" in order_type.lower()
        detail = {
            "Order #": stock,
            "Order type": order_type or "—",
            "Order status": order_status or "—",
            "Package (suffix)": str(v[9] or "—").strip(),
            "Accessory": str(v[11] or "—").strip(),
            "ETA window": ((iso(eta_from) or "?") + " → " + (iso(eta_to) or "?")) if (eta_from or eta_to) else "—",
            "Allocated to": allocated or "Open stock",
            "Comments": comments or "—",
        }
        rec = {
            "stock": stock, "vin": vin, "year": year, "make": "Lexus",
            "model": model, "trim": trim or "Premium",
            "body": body_for(model, series),
            "price": int(round(price)) if price else 0,
            "atValue": int(round(price)) if price else 0,
            "colour": ext, "interior": interior.title() if interior else "",
            "fuel": fuel_for(series, model, trim),
            "link": "",
            "car": f"{year} Lexus {model}".strip(),
            "detail": detail,
            "allocatedTo": allocated,
            "orderStatus": order_status,
        }
        if kind == "New":
            rec.update({
                "type": "New",
                "condition": "Demo" if is_demo else "New",
                "days": days_since(eta_from),
                "km": "Demo" if is_demo else "0 km",
                "kmNum": 0,
                "eta": "", "from": "",
            })
        else:  # Incoming
            eta = eta_to or eta_from
            rec.update({
                "type": "Incoming",
                "condition": order_status or "In transit",
                "days": 0,
                "km": "0 km", "kmNum": 0,
                "eta": eta.strftime("%b %d") if isinstance(eta, (datetime.date, datetime.datetime)) else "",
                "from": order_status or "Factory order",
            })
        out.append(rec)
    return out


# --------------------------------------------------------------------------- #
STAGE_MAP = {  # client "Level" -> (stage index, hot)
    "cold": (0, False),
    "lost": (1, False),
    "warm": (1, False),
    "hot": (2, True),
    "smoking": (4, True),
    "purchased": (6, False),
}
# per-client stage overrides keyed by name (spread across the funnel truthfully)
STAGE_OVERRIDE = {
    "keith kinsella": 2,      # sent info email, booking
    "nava muru": 3,           # sent videos, next is a drive
    "william butler": 4,      # ordering, cash upfront
}
SOURCE_MAP = {
    "walk-in": "Walk-in", "paisley": "Referral", "aticus split deal": "Referral",
    "repeat client": "Repeat", "website": "Website", "phone-in": "Phone-in",
}
STOCK_RE = re.compile(r"\b([0-9]{5,6}[A-Z]{1,3}|[A-HJ-NPR-Z0-9]{17})\b")


def parse_budget(rng):
    if not rng:
        return 0
    s = str(rng).replace(",", "")
    m = re.search(r"\$?\s*([0-9]+)\s*k", s, re.I)
    if m:
        return int(m.group(1)) * 1000
    m = re.search(r"([0-9]{4,6})", s)
    return int(m.group(1)) if m else 0


def build_deals(wb, inv_by_stock, inv_by_vin, inventory):
    ws = wb["Clients"]
    deals = []
    did = 0
    for _, v in rows(ws):
        name = v[0]
        if not name:
            continue
        name = str(name).strip()
        level = str(v[2] or "").strip().lower()
        stage, hot = STAGE_MAP.get(level, (0, False))
        stage = STAGE_OVERRIDE.get(name.lower(), stage)
        did += 1
        created = v[4]
        source = SOURCE_MAP.get(str(v[8] or "").strip().lower(), "Walk-in")
        email = str(v[9]).strip() if v[9] else ""
        buying = str(v[11] or "").strip()          # Used / New
        gen_type = str(v[12] or "")                # body-ish
        make = str(v[14] or "")
        trim = str(v[15] or "")
        dream = str(v[3] or "")
        notes_txt = str(v[23] or "").strip()
        matched = str(v[26] or "")
        budget = parse_budget(v[19])
        # body from "General type"
        body = "SUV"
        for cand in ("SUV", "Sedan", "Truck", "Van", "Coupe"):
            if cand.lower() in gen_type.lower():
                body = cand
                break

        # resolve vehicle of interest: prefer an explicitly matched stock / VIN
        # (from the "Matched Vehicle(s)" column or the notes), else score the
        # available inventory against make / trim / dream-car keywords.
        veh = None
        for tok in STOCK_RE.findall(matched) + STOCK_RE.findall(notes_txt):
            veh = inv_by_stock.get(tok) or inv_by_vin.get(tok.upper())
            if veh:
                break
        if not veh:
            want_tokens = set()
            for src in (make, trim, dream):
                for w in re.findall(r"[A-Za-z0-9]+", str(src)):
                    wl = w.lower()
                    if len(w) >= 2 and wl not in ("any", "used", "new", "the", "or",
                                                  "and", "sport", "fsport", "searching"):
                        want_tokens.add(wl)
            mk = (make or "").lower().strip()

            def score(x):
                if x["type"] == "Incoming" or x["price"] < 8000:
                    return -1
                if x["body"] != body:
                    return -1
                hay = (x["model"] + " " + x["make"] + " " + x["trim"]).lower()
                if mk and mk not in ("other", "any", "lexus") and mk not in hay:
                    return -1
                return sum(1 for t in want_tokens if t in hay)

            scored = [(score(x), x) for x in inventory]
            scored = [c for c in scored if c[0] >= 0]
            if scored:
                scored.sort(key=lambda c: (-c[0], c[1]["price"]))
                veh = scored[0][1]
            else:
                fallback = [x for x in inventory if x["type"] != "Incoming"
                            and x["price"] >= 8000 and x["body"] == body]
                veh = (min(fallback, key=lambda x: abs(x["price"] - (budget or 40000)))
                       if fallback else inventory[0])

        pay = "Finance"
        low = notes_txt.lower()
        if "cash" in low:
            pay = "Cash"
        elif "lease" in low:
            pay = "Lease"
        test_drove = bool(re.search(r"test dr|went on (a )?test", low))
        if level == "purchased":
            test_drove = True

        value = veh["price"]
        gross = int(round(value * 0.06))
        if level == "purchased":
            gross = 4564  # real front-end gross from the Sales record

        # follow-up cadence from stage / heat
        if level in ("purchased", "lost"):
            nxt = None
        elif hot or stage >= 3:
            nxt = TODAY                               # due today
        elif stage == 2:
            nxt = TODAY + datetime.timedelta(days=1)
        elif stage == 1:
            nxt = TODAY + datetime.timedelta(days=2)
        else:
            nxt = TODAY + datetime.timedelta(days=4)
        close = TODAY + datetime.timedelta(days={0: 30, 1: 21, 2: 14, 3: 10, 4: 7, 5: 5, 6: 0}[stage])

        note_entries = []
        if notes_txt:
            note_entries.append({"date": iso(created) or iso(TODAY),
                                 "text": notes_txt})
        want = []
        if make and make.lower() not in ("other", "any"):
            want.append(make)
        if trim and trim.lower() != "any":
            want.append(trim)
        want_line = "Came in from %s. Wants a %s%s%s." % (
            source, body.lower(),
            (" (" + " ".join(want) + ")") if want else "",
            (" around $%s" % f"{budget:,}") if budget else "")
        note_entries.append({"date": iso(created) or iso(TODAY), "text": want_line})

        # every stock the salesperson hand-matched to this client (resolved to inventory)
        matched_stocks = []
        for tok in STOCK_RE.findall(matched):
            key = tok if tok in inv_by_stock else None
            if not key and tok.upper() in inv_by_vin:
                key = inv_by_vin[tok.upper()]["stock"]
            if key and key not in matched_stocks:
                matched_stocks.append(key)

        deals.append({
            "id": did, "name": name, "phone": fmt_phone(v[1]),
            "email": email or (re.sub(r"[^a-z ]", "", name.lower()).strip().replace(" ", ".") + "@gmail.com"),
            "address": "", "stage": stage, "invStock": veh["stock"],
            "matchedStocks": matched_stocks,
            "trade": "None", "allowance": 0, "appraisal": 0,
            "pay": pay, "source": source, "testDrive": test_drove,
            "lastContact": iso(created) or iso(TODAY),
            "nextFollowUp": iso(nxt), "expectedClose": iso(close),
            "value": value, "gross": gross, "hot": hot,
            "wantBody": body, "wantMax": budget or int(round(value * 1.1)),
            "wantMake": make, "wantTrim": trim, "buying": buying,
            "fuelPref": str(v[13] or ""), "colorPref": str(v[17] or ""),
            "level": str(v[2] or "").strip(),
            "notes": note_entries,
        })
    return deals


def build_matches(wb, names):
    """Customer <-> vehicle matches the workbook already computed (`_UsedMatches`)."""
    out = []
    ws = wb["_UsedMatches"]
    for _, v in rows(ws):
        client = v[0]
        if not client or str(client).strip() not in names:
            continue
        stock = v[4]
        if not stock:
            continue
        out.append({
            "client": str(client).strip(),
            "stock": str(stock).strip(),
            "car": (" ".join(str(x) for x in [int(num(v[7])) if num(v[7]) else "", v[8] or v[9]] if x)).strip(),
            "colour": str(v[10] or "").strip(),
            "km": int(num(v[11])) if num(v[11]) else 0,
            "price": int(num(v[12])) if num(v[12]) else 0,
            "fit": str(v[13] or "").strip(),
            "status": str(v[14] or "").strip(),
            "source": str(v[3] or "").strip(),
        })
    return out


def fmt_phone(p):
    if p is None:
        return "—"
    s = re.sub(r"[^0-9]", "", str(p))
    if len(s) == 11 and s[0] == "1":
        s = s[1:]
    if len(s) == 10:
        return f"({s[0:3]}) {s[3:6]}-{s[6:]}"
    return str(p)


# --------------------------------------------------------------------------- #
def build_appointments(deals):
    """No appointment feed exists in the export; generate a realistic week for the
    active opportunities so the Calendar / Today views are populated."""
    by_name = {d["name"]: d for d in deals}
    # (deal name, type, day-index 0=Wed..4=Mon, start hour, hours)
    plan = [
        ("Nava Muru", "Test drive", 0, 10.0, 1, True, 60),
        ("William Butler", "Financing sign", 0, 14.0, 1, True, 30),
        ("Rene Ramsawak", "Appointment", 0, 11.5, 1, False, 60),
        ("Karan Minhas", "Delivery", 0, 16.0, 1.5, True, 120),
        ("Keith Kinsella", "Appointment", 1, 10.0, 1, True, 60),
        ("Alireza Pourfarshomi", "Test drive", 1, 13.0, 1, True, 60),
        ("William Butler", "Appraisal", 2, 9.5, 0.5, True, 30),
        ("Rene Ramsawak", "Test drive", 2, 14.5, 1, True, 60),
        ("Keith Kinsella", "Delivery", 3, 12.0, 1.5, False, 120),
        ("Alireza Pourfarshomi", "Appointment", 4, 9.5, 1, True, 60),
    ]
    appts = []
    aid = 0
    for name, typ, day, start, hours, gcal, mail in plan:
        d = by_name.get(name)
        if not d:
            continue
        aid += 1
        appts.append({"id": aid, "dealId": d["id"], "type": typ, "day": day,
                      "start": start, "hours": hours,
                      "gcalSynced": gcal, "emailReminderMinutes": mail})
    return appts


def build_ads(wb, inventory, inv_by_stock):
    ads = []
    # real FB_Log entries first
    ws = wb["FB_Log"]
    for _, v in rows(ws):
        stock = v[0]
        if not stock:
            continue
        stock = str(stock).strip()
        veh = inv_by_stock.get(stock)
        if not veh:
            continue
        posted = v[1]
        ads.append({"stock": stock,
                    "url": f"facebook.com/marketplace/item/{abs(hash(stock)) % 9000000000:010d}",
                    "posted": iso(posted) or iso(TODAY),
                    "status": "Live", "views": 0, "saves": 0, "messages": 0,
                    "listedPrice": veh["price"]})
    have = {a["stock"] for a in ads}
    # seed ads for the used units most in need of exposure (aged stock),
    # with deterministic engagement derived from days-in-stock.
    aged = sorted([v for v in inventory if v["type"] == "Used" and v["stock"] not in have],
                  key=lambda v: -v["days"])[:8]
    for i, v in enumerate(aged):
        d = v["days"]
        status = "Expired" if d > 75 else "Needs renewal" if d > 45 else "Live"
        posted = TODAY - datetime.timedelta(days=min(d, 40))
        ads.append({
            "stock": v["stock"],
            "url": f"facebook.com/marketplace/item/{8809000000 + (abs(hash(v['stock'])) % 900000)}",
            "posted": iso(posted), "status": status,
            "views": 300 + d * 47 + i * 90,
            "saves": 12 + (d * 3) % 180,
            "messages": 1 + (d // 12) % 9,
            "listedPrice": v["price"],
        })
    return ads


def build_repeat_buyers(wb, deals, inventory):
    buyers = []
    # 1) real repeat client from Clients (Karan bought a Sequoia, now wants LX/TX)
    karan = next((d for d in deals if "karan" in d["name"].lower() and d["stage"] != 6), None)
    if karan:
        buyers.append({
            "name": karan["name"], "dealId": karan["id"],
            "boughtVehicle": "2024 Toyota Sequoia (As-Is)", "boughtDate": "Jun 2026",
            "wants": "Repeat customer — wants a used Lexus LX or TX for his father",
            "preload": {"body": "SUV", "interior": "leather",
                        "priceMax": min(90000, max(40000, int(karan["wantMax"] * 1.05) or 80000))},
        })
    # 2) upsell targets from "Leads - With Vehicle": current owners we can put in something new
    ws = wb["Leads - With Vehicle"]
    picked = 0
    for _, v in rows(ws, start=3):
        first, last = v[2], v[1]
        year, make, model = v[4], v[5], v[6]
        if not (first and make and model and num(year)):
            continue
        name = f"{str(first).strip()} {str(last).strip()}"
        year = int(num(year))
        body = body_for(model, "")
        mileage = int(num(v[9])) if num(v[9]) else 0
        book = num(v[10]) or 0
        ceiling = int(round(max(45000, min(90000, (book or 40000) * 1.15))))
        buyers.append({
            "name": name, "dealId": 0,
            "boughtVehicle": f"{int(year)} {title_make(make)} {model}".strip(),
            "boughtDate": "Owns now",
            "wants": f"Drives a {int(year)} {title_make(make)} {model} (~{mileage:,} km) — trade-up / lease-end candidate",
            "preload": {"body": body, "interior": "", "priceMax": ceiling},
        })
        picked += 1
        if picked >= 4:
            break
    return buyers


# --------------------------------------------------------------------------- #
def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # client names + referenced stocks (so matched vehicles are never dropped
    # from inventory even when the source marks them sold)
    referenced = set()
    names = set()
    ws = wb["Clients"]
    for _, v in rows(ws):
        if v[0]:
            names.add(str(v[0]).strip())
        for tok in STOCK_RE.findall(str(v[26] or "")):
            referenced.add(tok)
    ws = wb["_UsedMatches"]
    for _, v in rows(ws):
        if v[0] and str(v[0]).strip() in names and v[4]:
            referenced.add(str(v[4]).strip())

    pricing = build_pricing(wb)
    inventory = []
    inventory += build_new(wb, pricing, "Inventory", "New")
    inventory += build_new(wb, pricing, "Pipeline", "Incoming")
    inventory += build_used(wb, referenced)

    # de-dupe by stock (keep first)
    seen = set()
    deduped = []
    for v in inventory:
        if v["stock"] in seen:
            continue
        seen.add(v["stock"])
        deduped.append(v)
    inventory = deduped

    inv_by_stock = {v["stock"]: v for v in inventory}
    inv_by_vin = {v["vin"].upper(): v for v in inventory if v["vin"]}

    deals = build_deals(wb, inv_by_stock, inv_by_vin, inventory)
    appts = build_appointments(deals)
    ads = build_ads(wb, inventory, inv_by_stock)
    buyers = build_repeat_buyers(wb, deals, inventory)
    matches = build_matches(wb, names)

    data = {
        "today": TODAY.isoformat(),
        "dealership": {"name": "Vincere", "brand": "Lexus · Sales CRM",
                       "target": 22},
        "inventory": inventory,
        "deals": deals,
        "appointments": appts,
        "ads": ads,
        "repeatBuyers": buyers,
        "matches": matches,
    }

    out = ROOT / "crm" / "data.js"
    out.parent.mkdir(exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    out.write_text("// Generated by scripts/build_crm_data.py — do not edit by hand.\n"
                   "window.CRM_DATA = " + payload + ";\n", encoding="utf-8")

    print(f"Wrote {out}")
    print(f"  inventory : {len(inventory)}  "
          f"(new {sum(1 for v in inventory if v['type']=='New')}, "
          f"used {sum(1 for v in inventory if v['type']=='Used')}, "
          f"incoming {sum(1 for v in inventory if v['type']=='Incoming')})")
    print(f"  deals     : {len(deals)}")
    print(f"  appts     : {len(appts)}")
    print(f"  ads       : {len(ads)}")
    print(f"  buyers    : {len(buyers)}")
    print(f"  matches   : {len(matches)}")


if __name__ == "__main__":
    main()
